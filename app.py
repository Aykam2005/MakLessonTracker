
from flask import Flask, session, redirect, url_for, request, render_template
import msal
import requests
from openpyxl import load_workbook
import os
from config import CLIENT_ID, CLIENT_SECRET, TENANT_ID, AUTHORITY, REDIRECT_URI, SCOPE

app = Flask(__name__)
app.secret_key = os.urandom(24)

def _build_msal_app():
    return msal.ConfidentialClientApplication(
        CLIENT_ID, authority=AUTHORITY,
        client_credential=CLIENT_SECRET
    )

def _build_auth_code_flow():
    return _build_msal_app().initiate_auth_code_flow(
        scopes=SCOPE,
        redirect_uri=REDIRECT_URI
    )

@app.route("/")
def index():
    if "token" not in session:
        return redirect(url_for("login"))
    return render_template("index.html")

@app.route("/login")
def login():
    session["flow"] = _build_auth_code_flow()
    return redirect(session["flow"]["auth_uri"])

@app.route("/callback")
def callback():
    flow = session.get("flow", {})
    result = _build_msal_app().acquire_token_by_authorization_code(
        request.args["code"],
        scopes=SCOPE,
        redirect_uri=REDIRECT_URI
    )
    if "access_token" in result:
        session["token"] = result["access_token"]
        return redirect(url_for("index"))
    return f"Login failed: {result.get('error_description')}"

@app.route("/read_excel")
def read_excel():
    if "token" not in session:
        return redirect(url_for("login"))
    
    file_path = get_excel_file_from_onedrive(session["token"], "LessonTracker.xlsm")
    if file_path:
        wb = load_workbook(file_path, keep_vba=True)
        sheet = wb["lesson log"]
        last_row = sheet.max_row
        last_entry = [cell.value for cell in sheet[last_row]]
        return f"Last lesson entry: {last_entry}"
    return "Failed to load Excel file."

def get_excel_file_from_onedrive(access_token, file_name):
    headers = {'Authorization': f'Bearer {access_token}'}
    url = f'https://graph.microsoft.com/v1.0/me/drive/root:/{file_name}:/content'

    response = requests.get(url, headers=headers)
    if response.status_code == 200:
        local_file = f"temp_{file_name}"
        with open(local_file, 'wb') as f:
            f.write(response.content)
        return local_file
    print(f"Graph API error: {response.status_code} - {response.text}")
    return None

if __name__ == "__main__":
    app.run(debug=True)
